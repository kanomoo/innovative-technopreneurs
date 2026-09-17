import os
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import win32com.client
import fitz

doc = docx.Document()
sec = doc.sections[0]
sec.top_margin = Inches(1.5)
sec.bottom_margin = Inches(1.0)
sec.left_margin = Inches(1.5)
sec.right_margin = Inches(1.0)

p_title = doc.add_paragraph()
p_title.paragraph_format.space_before = Pt(0)
p_title.paragraph_format.space_after = Pt(8)
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p_title.add_run('สารบัญ (ต่อ)')
r.font.name = 'TH Sarabun PSK'
r.font.size = Pt(18)
r.bold = True

items = [
    ('บทที่ 6  แบบจำลองผืนผ้าใบธุรกิจ (Business Model Canvas - BMC)', '19', True, 0),
    ('6.0  แผนภาพสรุป Business Model Canvas 9 ช่อง', '19', False, 0.25),
    ('6.1  พันธมิตรหลัก (Key Partners)', '19', False, 0.25),
    ('6.2  กิจกรรมหลัก (Key Activities)', '19', False, 0.25),
    ('6.3  ทรัพยากรหลัก (Key Resources)', '19', False, 0.25),
    ('6.4  คุณค่าที่ส่งมอบ (Value Propositions)', '19', False, 0.25),
    ('6.5  ความสัมพันธ์กับลูกค้า (Customer Relationships)', '19', False, 0.25),
    ('6.6  ช่องทางการเข้าถึง (Channels)', '21', False, 0.25),
    ('6.7  กลุ่มลูกค้าเป้าหมาย (Customer Segments)', '21', False, 0.25),
    ('6.8  โครงสร้างต้นทุน (Cost Structure)', '21', False, 0.25),
    ('6.9  แหล่งที่มารายได้ (Revenue Streams)', '21', False, 0.25),
    ('บทที่ 7  กลยุทธ์การตลาดและการขาย (Marketing & Sales Strategy)', '22', True, 0),
    ('7.1  การวางตำแหน่งทางการตลาด (Market Positioning Map)', '22', False, 0.25),
    ('7.2  กลยุทธ์ส่วนประสมทางการตลาด 4Ps สำหรับนวัตกรรมบริการ', '23', False, 0.25),
    ('7.3  กลยุทธ์การเจาะตลาดและการส่งเสริมการขาย (Go-to-Market)', '23', False, 0.25),
    ('บทที่ 8  กลยุทธ์ด้านเทคโนโลยีและนวัตกรรม (3 หมวด 10 ยุทธวิธี)', '24', True, 0),
    ('8.1  การวิเคราะห์ตามโมเดล Ten Types of Innovation', '24', False, 0.25),
    ('8.2  การคุ้มครองทรัพย์สินทางปัญญาและความได้เปรียบทางเทคโนโลยี', '25', False, 0.25),
    ('บทที่ 9  การประเมินและการบริหารจัดการความเสี่ยง (Risk Management)', '26', True, 0),
    ('9.1  การบ่งชี้ความเสี่ยง 6 ด้านตามบทเรียนที่ 11', '26', False, 0.25),
    ('9.2  เมทริกซ์การประเมินระดับความเสี่ยง (Risk Assessment Matrix)', '28', False, 0.25),
    ('9.3  แผนบริหารความต่อเนื่องทางธุรกิจ (Business Continuity Plan: BCP)', '28', False, 0.25),
    ('บทที่ 10  บทสรุป แผนการดำเนินงาน และความเป็นไปได้ทางการเงิน', '29', True, 0),
    ('10.1  บทสรุปโครงการ (Executive Summary)', '29', False, 0.25),
    ('10.2  แผนงานและเส้นทางการดำเนินโครงการ (Project Roadmap)', '29', False, 0.25),
    ('10.3  การวิเคราะห์จุดคุ้มทุนและความเป็นไปได้ทางการเงิน', '31', False, 0.25),
    ('10.4  ปัจจัยแห่งความสำเร็จ (Critical Success Factors)', '31', False, 0.25),
    ('10.5  ข้อเสนอแนะและทิศทางการพัฒนาในอนาคต', '31', False, 0.25),
    ('บรรณานุกรม (References)', '33', True, 0),
    ('ภาคผนวก ก  ภาพถ่ายโปสเตอร์ Business Model Canvas ต้นฉบับ', '34', True, 0),
    ('ภาคผนวก ข  สรุปผลการสำรวจความคิดเห็นกลุ่มตัวอย่างในโรงอาหาร มจพ.', '35', True, 0),
]

for title, pg, is_bold, ind in items:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    if ind > 0:
        p.paragraph_format.left_indent = Inches(ind)
    else:
        p.paragraph_format.left_indent = Inches(0)
    p.paragraph_format.first_line_indent = Inches(0)
    pPr = p._p.get_or_add_pPr()
    xml_str = f'<w:tabs {nsdecls("w")}><w:tab w:val="right" w:leader="dot" w:pos="8300"/></w:tabs>'
    pPr.append(parse_xml(xml_str))
    r1 = p.add_run(title)
    r1.font.name = 'TH Sarabun PSK'
    r1.font.size = Pt(13.5)
    r1.bold = is_bold
    r2 = p.add_run(f'\t{pg}')
    r2.font.name = 'TH Sarabun PSK'
    r2.font.size = Pt(13.5)
    r2.bold = is_bold

tpath = r'C:\Project\innovative-technopreneurs\test_toc_part2.docx'
doc.save(tpath)

word = win32com.client.DispatchEx('Word.Application')
word.Visible = False
wdoc = word.Documents.Open(tpath)
tpdf = r'C:\Project\innovative-technopreneurs\test_toc_part2.pdf'
wdoc.SaveAs(tpdf, FileFormat=17)
wdoc.Close(False)
word.Quit()

fdoc = fitz.open(tpdf)
print(f'Test TOC Part 2 Total Pages in PDF: {len(fdoc)}')
