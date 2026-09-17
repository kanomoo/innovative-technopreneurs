import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from pythainlp.tokenize import word_tokenize

sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = r'C:\Project\innovative-technopreneurs\Project_Smart_Queue_Food'
ASSETS_DIR = os.path.join(BASE_DIR, 'assets')
DOCX_OUTPUT = os.path.join(BASE_DIR, 'รายงานโครงงาน_Smart_Queue_Food_ฉบับสมบูรณ์.docx')
PDF_OUTPUT = os.path.join(BASE_DIR, 'รายงานโครงงาน_Smart_Queue_Food_ฉบับสมบูรณ์.pdf')

FONT_NAME = 'TH Sarabun PSK'
CLR_BLACK = RGBColor(0, 0, 0)

def thai_zwsp(text):
    """
    Inserts zero-width spaces (\u200b) at word boundaries to allow natural
    word-wrapping in Word without breaking syllables or stretching characters.
    """
    if not text:
        return ""
    
    lines = str(text).split('\n')
    processed_lines = []
    for line in lines:
        if not line:
            processed_lines.append("")
            continue
        tokens = word_tokenize(line, engine='newmm')
        processed_lines.append('\u200b'.join(tokens))
    return '\n'.join(processed_lines)

def set_run_font(run, size_pt=16, bold=False, italic=False, underline=False):
    run.font.name = FONT_NAME
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.underline = underline
    run.font.color.rgb = CLR_BLACK
    
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:ascii'), FONT_NAME)
    rFonts.set(qn('w:hAnsi'), FONT_NAME)
    rFonts.set(qn('w:cs'), FONT_NAME)
    rFonts.set(qn('w:eastAsia'), FONT_NAME)
    
    # Language tag
    lang = parse_xml(f'<w:lang {nsdecls("w")} w:val="th-TH" w:eastAsia="th-TH" w:bidi="th-TH"/>')
    rPr.append(lang)
    
    # Disable proofing (no red/blue squiggly lines)
    no_proof = parse_xml(f'<w:noProof {nsdecls("w")}/>')
    rPr.append(no_proof)

def format_paragraph(p, space_before=0, space_after=3, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.THAI_JUSTIFY, keep_with_next=False):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    p.alignment = align
    p.paragraph_format.widow_control = True
    if keep_with_next:
        p.paragraph_format.keep_with_next = True

def add_body_p(doc, text, bold_prefix=None, indent=True, space_after=3, page_break_before=False):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=0, space_after=space_after, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.THAI_JUSTIFY)
    if indent:
        # Standard Thai academic indent: 0.5 inches (1.27 cm / 1 tab)
        p.paragraph_format.first_line_indent = Inches(0.5)
    else:
        p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.left_indent = Inches(0)
    
    if page_break_before:
        p.paragraph_format.page_break_before = True
    
    if bold_prefix:
        r_bold = p.add_run(thai_zwsp(bold_prefix))
        set_run_font(r_bold, size_pt=16, bold=True)
    
    r_text = p.add_run(thai_zwsp(text))
    set_run_font(r_text, size_pt=16, bold=False)
    return p

def add_numbered_item(doc, num_label, text, bold_prefix=None, space_after=2, keep_with_next=False, page_break_before=False):
    """
    Standard Thai academic indented numbered list.
    left_indent = 0.5 in, first_line_indent = 0 in.
    Alignd cleanly at 0.5 in with no weird stretched spacing.
    """
    p = doc.add_paragraph()
    format_paragraph(p, space_before=0, space_after=space_after, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT, keep_with_next=keep_with_next)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(0)
    if page_break_before:
        p.paragraph_format.page_break_before = True
    
    # Number label (e.g. "1)  ")
    r_num = p.add_run(f"{num_label}  ")
    set_run_font(r_num, size_pt=16, bold=False)
    
    if bold_prefix:
        r_bold = p.add_run(thai_zwsp(bold_prefix))
        set_run_font(r_bold, size_pt=16, bold=True)
    
    r_text = p.add_run(thai_zwsp(text))
    set_run_font(r_text, size_pt=16, bold=False)
    return p

def add_bullet_item(doc, text, bold_prefix=None, space_after=2, keep_with_next=False, page_break_before=False):
    """
    Standard Thai academic indented bullet item.
    left_indent = 0.5 in, first_line_indent = 0 in.
    """
    p = doc.add_paragraph()
    format_paragraph(p, space_before=0, space_after=space_after, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT, keep_with_next=keep_with_next)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(0)
    if page_break_before:
        p.paragraph_format.page_break_before = True
    
    r_sym = p.add_run("•  ")
    set_run_font(r_sym, size_pt=16, bold=True)
    
    if bold_prefix:
        r_bold = p.add_run(thai_zwsp(bold_prefix))
        set_run_font(r_bold, size_pt=16, bold=True)
    
    r_text = p.add_run(thai_zwsp(text))
    set_run_font(r_text, size_pt=16, bold=False)
    return p

def add_heading_1(doc, text, page_break_before=False):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=12, space_after=3, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT, keep_with_next=True)
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.left_indent = Inches(0)
    if page_break_before:
        p.paragraph_format.page_break_before = True
    r = p.add_run(thai_zwsp(text))
    # Standard KMUTNB thesis/project heading 1 is 16pt bold (flush left)
    set_run_font(r, size_pt=16, bold=True)
    return p

def add_heading_2(doc, text, page_break_before=False):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=8, space_after=2, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT, keep_with_next=True)
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.left_indent = Inches(0)
    if page_break_before:
        p.paragraph_format.page_break_before = True
    r = p.add_run(thai_zwsp(text))
    set_run_font(r, size_pt=16, bold=True)
    return p

def add_heading_3(doc, text, page_break_before=False):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=6, space_after=2, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT, keep_with_next=True)
    p.paragraph_format.first_line_indent = Inches(0.5)
    p.paragraph_format.left_indent = Inches(0)
    if page_break_before:
        p.paragraph_format.page_break_before = True
    r = p.add_run(thai_zwsp(text))
    set_run_font(r, size_pt=16, bold=True)
    return p

def add_chapter_title(doc, chapter_num, chapter_title):
    p1 = doc.add_paragraph()
    format_paragraph(p1, space_before=14, space_after=3, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.CENTER, keep_with_next=True)
    p1.paragraph_format.first_line_indent = Inches(0)
    r1 = p1.add_run(f"บทที่ {chapter_num}")
    set_run_font(r1, size_pt=20, bold=True)
    
    p2 = doc.add_paragraph()
    format_paragraph(p2, space_before=0, space_after=14, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.CENTER, keep_with_next=True)
    p2.paragraph_format.first_line_indent = Inches(0)
    r2 = p2.add_run(thai_zwsp(chapter_title))
    set_run_font(r2, size_pt=18, bold=True)
    return p2

def add_figure(doc, image_path, figure_num, caption_text, width_inches=4.2):
    if not os.path.exists(image_path):
        raise FileNotFoundError(f"CRITICAL ERROR: Image not found for Figure {figure_num}: {image_path}")
    p_img = doc.add_paragraph()
    format_paragraph(p_img, space_before=6, space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER, keep_with_next=True)
    p_img.paragraph_format.first_line_indent = Inches(0)
    run_img = p_img.add_run()
    run_img.add_picture(image_path, width=Inches(width_inches))
    
    p_cap = doc.add_paragraph()
    format_paragraph(p_cap, space_before=2, space_after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    p_cap.paragraph_format.first_line_indent = Inches(0)
    
    r_num = p_cap.add_run(f"รูปที่ {figure_num}  ")
    set_run_font(r_num, size_pt=15, bold=True)
    
    r_cap = p_cap.add_run(thai_zwsp(caption_text))
    set_run_font(r_cap, size_pt=15, bold=False)
    return p_cap

def add_styled_table(doc, table_num, caption_text, headers, rows_data, col_widths=None, page_break_before=False, font_size_pt=14, padding_twips=50):
    # Caption with keep_with_next to NEVER separate from table
    p_cap = doc.add_paragraph()
    format_paragraph(p_cap, space_before=8, space_after=3, align=WD_ALIGN_PARAGRAPH.LEFT, keep_with_next=True)
    p_cap.paragraph_format.first_line_indent = Inches(0)
    if page_break_before:
        p_cap.paragraph_format.page_break_before = True
        
    r_num = p_cap.add_run(f"ตารางที่ {table_num}  ")
    set_run_font(r_num, size_pt=15, bold=True)
    r_cap = p_cap.add_run(thai_zwsp(caption_text))
    set_run_font(r_cap, size_pt=15, bold=False)
    
    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Strict Academic Table Borders (Top single, bottom single, no vertical)
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="12" w:space="0" w:color="000000"/>'
        f'<w:bottom w:val="single" w:sz="12" w:space="0" w:color="000000"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="D0D0D0"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)
    
    # Header row
    hdr_row = table.rows[0]
    trPr_hdr = hdr_row._tr.get_or_add_trPr()
    # Repeat header row on every page & prevent split
    trPr_hdr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    trPr_hdr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    
    hdr_cells = hdr_row.cells
    for col_idx, text in enumerate(headers):
        cell = hdr_cells[col_idx]
        tcPr = cell._tc.get_or_add_tcPr()
        if col_widths and col_idx < len(col_widths):
            w_in = col_widths[col_idx]
            w_float = w_in.inches if hasattr(w_in, 'inches') else float(w_in)
            twips = int(w_float * 1440)
            tcW = parse_xml(f'<w:tcW {nsdecls("w")} w:w="{twips}" w:type="dxa"/>')
            tcPr.append(tcW)
            
        tcBorders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(tcBorders)
        
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{padding_twips}" w:type="dxa"/><w:bottom w:w="{padding_twips}" w:type="dxa"/><w:left w:w="100" w:type="dxa"/><w:right w:w="100" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)
        
        p = cell.paragraphs[0]
        format_paragraph(p, space_before=0, space_after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
        p.paragraph_format.first_line_indent = Inches(0)
        r = p.add_run(thai_zwsp(text))
        set_run_font(r, size_pt=font_size_pt, bold=True)
    
    # Data rows
    for row_idx, r_data in enumerate(rows_data):
        row = table.rows[row_idx + 1]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        
        row_cells = row.cells
        for col_idx, text in enumerate(r_data):
            cell = row_cells[col_idx]
            tcPr = cell._tc.get_or_add_tcPr()
            if col_widths and col_idx < len(col_widths):
                w_in = col_widths[col_idx]
                w_float = w_in.inches if hasattr(w_in, 'inches') else float(w_in)
                twips = int(w_float * 1440)
                tcW = parse_xml(f'<w:tcW {nsdecls("w")} w:w="{twips}" w:type="dxa"/>')
                tcPr.append(tcW)
            tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{padding_twips}" w:type="dxa"/><w:bottom w:w="{padding_twips}" w:type="dxa"/><w:left w:w="100" w:type="dxa"/><w:right w:w="100" w:type="dxa"/></w:tcMar>')
            tcPr.append(tcMar)
            
            p = cell.paragraphs[0]
            format_paragraph(p, space_before=0, space_after=0, align=WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.first_line_indent = Inches(0)
            r = p.add_run(thai_zwsp(str(text)))
            set_run_font(r, size_pt=font_size_pt, bold=False)

def add_reference_item(doc, ref_text):
    """
    Standard APA Hanging Indent Reference Item:
    First line: flush left (0 inches)
    Subsequent lines: indented 0.5 inches
    """
    p = doc.add_paragraph()
    format_paragraph(p, space_before=2, space_after=4, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)
    r = p.add_run(thai_zwsp(ref_text))
    set_run_font(r, size_pt=15, bold=False)
    return p

print('Upgraded Academic report helpers loaded successfully.')
