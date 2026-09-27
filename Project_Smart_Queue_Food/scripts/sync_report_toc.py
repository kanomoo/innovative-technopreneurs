import json,re,sys
from pathlib import Path
from docx import Document
from copy import deepcopy
root=Path(__file__).resolve().parents[2]/'scratch'/'report-reconciliation'
data=json.loads((root/'pdf-page-map.json').read_text(encoding='utf8'))
d=Document(root/'reconciled-report.docx')
start=data['body_start_physical']-1
def norm(t):return re.sub(r'\s+','',t).replace('\u200b','').replace('ำ','า')
verified=[]
for p in d.paragraphs[:98]:
 if '\t' not in p.text:continue
 label,old=p.text.rsplit('\t',1)
 key=norm(label)
 if not old.strip().isdigit():continue
 m=re.match(r'(บทที่\d+|\d+\.\d+|ตารางที่(?:\d+|ข)\.\d+|รูปที่(?:\d+|ก)\.\d+|บรรณานุกรม|ภาคผนวก[กข])',key)
 if not m:raise ValueError(label)
 key=m.group(1)
 # A chapter key must not match the prefix of chapter 10.
 if key.startswith(('บทที่','ภาคผนวก','บรรณานุกรม')):
  candidates=[i for i,lines in enumerate(data['lines']) if i>=start and key in lines]
 else:
  matches=[q for q in d.paragraphs[98:] if norm(q.text).startswith(key) and (len(norm(q.text))==len(key) or not norm(q.text)[len(key)].isdigit())]
  if key.startswith(('รูปที่','ตารางที่')):
   matches=[q for q in matches if q.runs and q.runs[0].font.size and q.runs[0].font.size.pt<16]
  if not matches:raise ValueError('No heading/caption: '+key)
  query=norm(matches[0].text)[:45]
  candidates=[i for i,t in enumerate(data['text']) if i>=start and query in t]
 if not candidates:raise ValueError('No page: '+key)
 page=candidates[0]-start+1
 if key=='4.3':label='4.3  ภาพจำลองส่วนต่อประสานผู้ใช้ (UI Concept)'
 if key=='รูปที่4.2':label='รูปที่ 4.2  ภาพจำลองส่วนต่อประสานผู้ใช้เพื่ออธิบายแนวคิด'
 rp=deepcopy(p.runs[0]._r.rPr) if p.runs else None
 p.clear();r=p.add_run(label+'\t'+str(page))
 if rp is not None:r._r.insert(0,rp)
 verified.append({'key':key,'page':page,'physical_page':candidates[0]+1})
d.save(root/'final-report.docx')
(root/'toc-verification.json').write_text(json.dumps(verified,ensure_ascii=False,indent=2),encoding='utf8')
print('Updated',len(verified),'entries')
